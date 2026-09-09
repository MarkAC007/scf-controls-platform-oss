"""Tests for the Word (.docx) export.

Two halves. The first is pure: ``render_docx`` is a function from Markdown to
bytes, so the assertions open the package it produced and look at what Word
would actually see -- paragraph styles, run formatting, table objects -- rather
than at an intermediate string. Asserting on the HTML in between would pass
happily while the Word file came out empty.

The second half exercises the export endpoint over ASGI, because two of the
things that matter about this feature are properties of the route and not of
the renderer: that ``?format=docx`` is accepted and ``?format=rtf`` is not, and
that the Word branch returns before the endpoint ever reads the organisation
record. That ordering is the no-branding guarantee, and the only way to test a
guarantee about code that is *not* reached is to make reaching it fail.

``dependency_overrides`` cannot reach the ``require_org_role`` closures --
FastAPI 0.141 hides included routes behind ``_IncludedRouter`` -- so auth is
stubbed at the module the closure resolves through, the way
test_team_assignment_list_filters.py does it.
"""
import io
import zipfile
from contextlib import asynccontextmanager
from types import SimpleNamespace
from uuid import uuid4

import docx
import httpx
import pytest

from services.doc_gen.renderer import render_docx

WORD_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

#: The organisation whose name must never appear in a Word export. Named after
#: a real tenant on purpose: the failure this guards against is a branded PDF
#: helper being reused for the .docx path, and that would put this exact string
#: into the file.
ORG_NAME = "Odin Medical Ltd"

PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"0" * 64


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _open(data: bytes) -> docx.document.Document:
    """Open rendered bytes as Word would. Raises if the package is malformed."""
    return docx.Document(io.BytesIO(data))


def _styles(document) -> list:
    return [(p.style.name, p.text) for p in document.paragraphs]


def _style_of(document, text: str):
    """Style name of the first paragraph whose text matches exactly."""
    for paragraph in document.paragraphs:
        if paragraph.text == text:
            return paragraph.style.name
    raise AssertionError(f"no paragraph reads {text!r}; got {_styles(document)}")


def _runs(document) -> list:
    return [r for p in document.paragraphs for r in p.runs]


def _run_named(document, text: str):
    for run in _runs(document):
        if run.text == text:
            return run
    raise AssertionError(f"no run reads {text!r}")


def _all_text(document) -> str:
    """Every string a reader could see, tables included.

    ``document.paragraphs`` is top level only, so a value that ended up in a
    table cell would be invisible to an assertion that only walked paragraphs.
    """
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _package_contains(data: bytes, needle: str) -> bool:
    """Search every decompressed part of the .docx for a string.

    Stronger than reading the extracted text: a branded template would put the
    organisation name in a header, a footer or the core properties, none of
    which reach ``document.paragraphs``. A raw ``in data`` check would find
    nothing at all -- the parts are deflated.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        for name in package.namelist():
            if needle.encode("utf-8") in package.read(name):
                return True
    return False


# ---------------------------------------------------------------------------
# The package itself
# ---------------------------------------------------------------------------


def test_the_export_is_an_ooxml_container():
    """A .docx is a zip. If the first two bytes are not PK, Word shows the user
    a repair dialogue rather than a document."""
    assert render_docx("# Policy\n\nBody.").startswith(b"PK")


def test_the_package_opens_as_a_word_document():
    document = _open(render_docx("# Policy\n\nBody."))
    assert "Body." in _all_text(document)


@pytest.mark.parametrize("source", ["", "   ", "\n\n", "<!-- CONFLICT: only a marker -->"])
def test_content_with_no_blocks_still_produces_an_openable_document(source):
    """A document whose operative content is empty -- a stub, or a section that
    held nothing but a merge marker -- must download as an empty Word file, not
    as a 500 and not as zero bytes."""
    data = render_docx(source)
    assert data.startswith(b"PK")
    _open(data)


def test_a_title_alone_still_produces_the_heading():
    document = _open(render_docx("", title="Access Control Policy"))
    assert _style_of(document, "Access Control Policy") == "Heading 1"


# ---------------------------------------------------------------------------
# Headings
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("level", [1, 2, 3, 4, 5, 6])
def test_headings_land_on_the_matching_word_heading_style(level):
    source = f"{'#' * level} Heading Text\n\nBody.\n"
    document = _open(render_docx(source))
    assert _style_of(document, "Heading Text") == f"Heading {level}"


def test_a_deep_heading_does_not_fail_the_export():
    """python-docx's default template defines Heading 1-9 today, but the style
    lookup degrades rather than raising precisely so that a template change --
    or a future custom template -- turns an h6 into a shallower heading instead
    of into a failed download."""
    document = _open(render_docx("###### Deepest\n"))
    assert _style_of(document, "Deepest").startswith("Heading")


def test_the_title_is_not_repeated_when_the_body_already_opens_with_one():
    document = _open(render_docx("# Access Policy\n\nBody.\n", title="Access Policy"))
    assert [t for _, t in _styles(document)].count("Access Policy") == 1


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------


def test_bold_lands_as_a_bold_run():
    document = _open(render_docx("Some **emphatic** text.\n"))
    assert _run_named(document, "emphatic").bold is True


def test_italic_lands_as_an_italic_run():
    document = _open(render_docx("Some *stressed* text.\n"))
    assert _run_named(document, "stressed").italic is True


def test_nested_emphasis_produces_one_run_that_is_both():
    """``***x***`` nests the tags. A walker that tracked a single current
    format would emit the inner tag only and lose the outer one."""
    run = _run_named(_open(render_docx("A ***loud*** word.\n")), "loud")
    assert (run.bold, run.italic) == (True, True)


def test_inline_code_lands_in_a_monospace_run():
    run = _run_named(_open(render_docx("Set `MAX_RETRIES` to 3.\n")), "MAX_RETRIES")
    assert run.font.name == "Courier New"


def test_bold_does_not_leak_into_the_text_that_follows_it():
    """The formatting of a tag's tail belongs to the parent, not the tag."""
    document = _open(render_docx("**bold** then plain.\n"))
    assert _run_named(document, " then plain.").bold is not True


def test_an_escaped_ampersand_reaches_the_document_as_one_character():
    """Content is HTML-escaped on its way through the shared Markdown parser.
    A reader must see ``A & B``, not ``A &amp; B``."""
    assert "Vendors A & B." in _all_text(_open(render_docx("Vendors A & B.\n")))


# ---------------------------------------------------------------------------
# Lists
# ---------------------------------------------------------------------------


def test_bullet_items_land_on_a_bullet_style():
    document = _open(render_docx("- alpha\n- beta\n"))
    assert _style_of(document, "alpha") == "List Bullet"
    assert _style_of(document, "beta") == "List Bullet"


def test_numbered_items_land_on_a_numbered_style():
    document = _open(render_docx("1. first\n2. second\n"))
    assert _style_of(document, "first") == "List Number"


def test_a_nested_list_steps_to_the_next_style():
    source = "- outer\n    - inner\n"
    document = _open(render_docx(source))
    assert _style_of(document, "outer") == "List Bullet"
    assert _style_of(document, "inner") == "List Bullet 2"


def test_a_list_nested_past_the_available_styles_still_indents():
    """python-docx's template stops at 'List Bullet 3'. A four-deep list has to
    keep stepping in on indentation rather than flattening or raising."""
    source = "- one\n    - two\n        - three\n            - four\n"
    document = _open(render_docx(source))
    deepest = next(p for p in document.paragraphs if p.text == "four")
    third = next(p for p in document.paragraphs if p.text == "three")
    assert deepest.style.name.startswith("List Bullet")
    assert (deepest.paragraph_format.left_indent or 0) > (
        third.paragraph_format.left_indent or 0
    )


def test_a_list_item_keeps_its_inline_formatting():
    document = _open(render_docx("- an **important** item\n"))
    assert _run_named(document, "important").bold is True


# ---------------------------------------------------------------------------
# Code blocks
# ---------------------------------------------------------------------------


CODE_MD = """```python
def check(value):
    return value

# after a blank line
```
"""


def test_a_fenced_code_block_keeps_its_line_breaks():
    """Every line has to remain separately addressable. A single paragraph with
    soft breaks looks the same in Word but reads back as one unbroken line."""
    text = [t for _, t in _styles(_open(render_docx(CODE_MD)))]
    assert "def check(value):" in text
    assert "    return value" in text
    assert "# after a blank line" in text


def test_a_blank_line_inside_a_code_block_survives():
    document = _open(render_docx(CODE_MD))
    lines = [p.text for p in document.paragraphs if p.style.name == "No Spacing"]
    assert "" in lines, lines


def test_code_block_lines_are_monospace():
    document = _open(render_docx(CODE_MD))
    run = _run_named(document, "def check(value):")
    assert run.font.name == "Courier New"


def test_the_closing_fence_does_not_add_a_trailing_blank_line():
    """``<pre><code>...\\n</code></pre>`` always carries a closing newline that
    is fence syntax, not the author's code."""
    lines = [
        p.text for p in _open(render_docx(CODE_MD)).paragraphs
        if p.style.name == "No Spacing"
    ]
    assert lines[-1] == "# after a blank line"


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


TABLE_MD = """| Control | Status | Owner |
|---------|--------|-------|
| AAA-01 | Implemented | Security Manager |
| AAA-02 | In Progress |
"""


def test_a_markdown_table_becomes_a_word_table():
    document = _open(render_docx(TABLE_MD))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 3)


def test_a_table_does_not_arrive_as_literal_pipes():
    """The failure this guards against is the whole table falling through the
    walker as paragraph text, which still 'exports' and is still unreadable."""
    document = _open(render_docx(TABLE_MD))
    assert "|" not in "\n".join(p.text for p in document.paragraphs)


def test_the_table_header_row_is_bold():
    cell = _open(render_docx(TABLE_MD)).tables[0].rows[0].cells[0]
    assert [run.bold for run in cell.paragraphs[0].runs] == [True]


def test_table_body_cells_are_not_bold():
    cell = _open(render_docx(TABLE_MD)).tables[0].rows[1].cells[0]
    assert all(run.bold is not True for run in cell.paragraphs[0].runs)


def test_a_ragged_row_leaves_the_missing_cell_empty_rather_than_raising():
    """A section a human edited can easily lose a pipe. One malformed row must
    not take the whole export down with it."""
    table = _open(render_docx(TABLE_MD)).tables[0]
    assert [c.text for c in table.rows[2].cells] == ["AAA-02", "In Progress", ""]


def test_the_table_carries_a_bordered_style():
    """Without borders the cells render as invisible boxes, which reads as
    mangled prose rather than as a table."""
    assert _open(render_docx(TABLE_MD)).tables[0].style.name == "Table Grid"


# ---------------------------------------------------------------------------
# Links, rules, quotes, and markup we do not map
# ---------------------------------------------------------------------------


def test_a_link_becomes_a_real_word_hyperlink():
    document = _open(render_docx("See [the policy](https://example.com/p) now.\n"))
    links = [(h.text, h.address) for p in document.paragraphs for h in p.hyperlinks]
    assert links == [("the policy", "https://example.com/p")]


def test_link_text_is_still_part_of_the_readable_document():
    document = _open(render_docx("See [the policy](https://example.com/p) now.\n"))
    assert "See the policy now." in _all_text(document)


def test_a_link_with_no_destination_stays_plain_text():
    """An empty relationship target produces a package Word refuses to open."""
    document = _open(render_docx("An [empty]() link.\n"))
    assert "empty" in _all_text(document)
    assert not [h for p in document.paragraphs for h in p.hyperlinks]


def test_a_link_with_no_text_falls_back_to_showing_the_url():
    """Otherwise the destination vanishes entirely from a printed document."""
    document = _open(render_docx("[](https://example.com/dest)\n"))
    assert "https://example.com/dest" in _all_text(document)


def test_a_horizontal_rule_draws_a_border():
    data = render_docx("Above.\n\n---\n\nBelow.\n")
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert "w:pBdr" in package.read("word/document.xml").decode()


def test_a_blockquote_is_indented():
    document = _open(render_docx("> Quoted guidance.\n"))
    quoted = next(p for p in document.paragraphs if "Quoted guidance." in p.text)
    assert (quoted.paragraph_format.left_indent or 0) > 0


def test_unmapped_markup_keeps_its_text():
    """An element with no mapping recurses into its children. Silently dropping
    the author's words is the one outcome that is worse than ugly output."""
    document = _open(render_docx("<dl><dd>Definition body.</dd></dl>\n"))
    assert "Definition body." in _all_text(document)


def test_merge_markers_never_reach_the_word_file():
    """Markers are review scaffolding, stripped in every export because an
    exported document is what an auditor reads."""
    source = "## Scope\n\n<!-- CONFLICT: edited and regenerated. -->\n\nBody text.\n"
    document = _open(render_docx(source))
    text = _all_text(document)
    assert "CONFLICT" not in text
    assert "Body text." in text


# ---------------------------------------------------------------------------
# No branding
#
# The .docx is handed over to be edited. render_pdf next door is branded on
# purpose -- masthead, logo, organisation name, footer -- and the risk this
# section exists for is somebody "finishing" the Word path by copying it.
# ---------------------------------------------------------------------------


BRANDED_LOOKING_MD = """# Information Security Policy

Issued under the ISMS. Owner: the CISO function.

| Field | Value |
|-------|-------|
| Version | 2.1 |
"""


def test_render_docx_cannot_be_asked_to_brand():
    """The guarantee is structural, not conditional. A flag can be passed
    wrongly; an argument that does not exist cannot be."""
    import inspect

    parameters = set(inspect.signature(render_docx).parameters)
    assert parameters == {"content", "title"}


def test_the_word_export_contains_no_images():
    data = render_docx(BRANDED_LOOKING_MD, title="Information Security Policy")
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert [n for n in package.namelist() if n.startswith("word/media/")] == []
    assert len(_open(data).inline_shapes) == 0


def test_the_word_export_never_names_the_organisation():
    """The renderer is never given the name, so the only way it could appear is
    a branded helper creeping into this path."""
    data = render_docx(BRANDED_LOOKING_MD, title="Information Security Policy")
    assert not _package_contains(data, ORG_NAME)


def test_organisation_text_the_author_typed_is_still_kept():
    """The rule is that the *platform* does not add branding, not that the word
    is banned -- a policy that names its own organisation in its prose must
    still export that sentence."""
    data = render_docx(f"This policy is issued by {ORG_NAME}.\n")
    assert ORG_NAME in _all_text(_open(data))


# ---------------------------------------------------------------------------
# The export endpoint
# ---------------------------------------------------------------------------


DOC_MD = "# Access Control Policy\n\n- one\n- two\n\nBody text.\n"


def _document():
    return SimpleNamespace(
        id=uuid4(),
        title="Access Control Policy",
        domain_id="IAC",
        merged_content=DOC_MD,
        lifecycle_status="published",
        generator_name="policy",
        document_type="policy",
        generation_version=2,
    )


class _RefusingSession:
    """A session that fails if the handler queries anything.

    ``_load_document`` is stubbed out, so the only query left in the export
    handler is the organisation-and-logo lookup that feeds the PDF masthead.
    Making it raise turns "the Word branch returns before the branding data is
    fetched" from a comment into an assertion.
    """

    def __init__(self):
        self.queried = False

    async def execute(self, *args, **kwargs):
        self.queried = True
        raise AssertionError("the docx branch must not read the organisation record")


class _OrgSession:
    """A session that would happily hand over a branded organisation record."""

    async def execute(self, *args, **kwargs):
        row = SimpleNamespace(
            name=ORG_NAME, logo_data=PNG_BYTES, logo_content_type="image/png"
        )
        return SimpleNamespace(one_or_none=lambda: row, scalar_one_or_none=lambda: row)


@asynccontextmanager
async def _client(session, document=None):
    import auth as auth_mod
    import main
    from api import documents as documents_api
    from database import get_db

    org_id = uuid4()
    user = SimpleNamespace(db_id=str(uuid4()), email="viewer@example.com",
                           auth_method="google")
    membership = SimpleNamespace(user=user, organization_id=org_id, role="viewer")
    loaded = document if document is not None else _document()

    async def _override_db():
        yield session

    async def _require_auth(*args, **kwargs):
        return user

    async def _verify(org, user_, db_, min_role="viewer"):
        return membership

    async def _load(db, document_id, organization_id, **kwargs):
        return loaded

    originals = (auth_mod.require_auth, auth_mod.verify_org_membership,
                 documents_api._load_document)
    auth_mod.require_auth = _require_auth
    auth_mod.verify_org_membership = _verify
    documents_api._load_document = _load
    main.app.dependency_overrides[get_db] = _override_db
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=main.app),
            base_url="http://export",
            headers={"Authorization": "Bearer stub"},
        ) as client:
            yield client, org_id, loaded
    finally:
        (auth_mod.require_auth, auth_mod.verify_org_membership,
         documents_api._load_document) = originals
        main.app.dependency_overrides.pop(get_db, None)


def _url(org_id, document, fmt):
    return f"/api/organizations/{org_id}/documents/{document.id}/export?format={fmt}"


@pytest.mark.asyncio
async def test_the_endpoint_returns_a_word_document():
    async with _client(_RefusingSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "docx"))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(WORD_MEDIA_TYPE)
    assert response.content.startswith(b"PK")


@pytest.mark.asyncio
async def test_the_endpoint_names_the_download_with_a_docx_extension():
    async with _client(_RefusingSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "docx"))

    disposition = response.headers["content-disposition"]
    assert disposition == 'attachment; filename="access-control-policy-iac.docx"'


@pytest.mark.asyncio
async def test_the_downloaded_bytes_are_the_document_a_reader_would_see():
    async with _client(_RefusingSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "docx"))

    assert "Body text." in _all_text(_open(response.content))


@pytest.mark.asyncio
async def test_the_docx_branch_never_reaches_the_organisation_record():
    """The branch sits ahead of the organisation lookup, and that placement is
    the guarantee rather than a tidy-up. This fails the moment it moves."""
    session = _RefusingSession()
    async with _client(session) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "docx"))

    assert response.status_code == 200
    assert session.queried is False


@pytest.mark.asyncio
async def test_a_branded_organisation_record_still_yields_an_unbranded_file():
    """Belt and braces for the test above: even handed an organisation with a
    name and a logo, the Word file carries neither."""
    async with _client(_OrgSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "docx"))

    assert response.status_code == 200
    assert not _package_contains(response.content, ORG_NAME)
    with zipfile.ZipFile(io.BytesIO(response.content)) as package:
        assert [n for n in package.namelist() if n.startswith("word/media/")] == []


@pytest.mark.asyncio
async def test_an_unknown_export_format_is_still_rejected():
    """Widening the pattern for docx must not widen it for anything else."""
    async with _client(_RefusingSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, "rtf"))

    assert response.status_code == 422


@pytest.mark.parametrize("fmt,media_type", [
    ("md", "text/markdown"),
    ("html", "text/html"),
])
@pytest.mark.asyncio
async def test_the_existing_text_exports_are_unchanged(fmt, media_type):
    """The docx branch was inserted between them and the PDF path."""
    async with _client(_RefusingSession()) as (client, org_id, document):
        response = await client.get(_url(org_id, document, fmt))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(media_type)
    assert "Body text." in response.text


# ---------------------------------------------------------------------------
# Style degradation
#
# python-docx's default template defines every style this renderer asks for, so
# none of the fallbacks below fire in production today. That is exactly why they
# need tests: an untested fallback is indistinguishable from a broken one, and
# the next person to read the guard will be tempted to delete it as dead code.
# The template is a library asset that changes between releases, and it is also
# what a future custom template would replace.
# ---------------------------------------------------------------------------


class _PartialTemplate:
    """A style table missing the names it was told to withhold."""

    def __init__(self, missing):
        self._missing = set(missing)

    def __getitem__(self, name):
        if name in self._missing:
            raise KeyError(name)
        return object()


def test_style_resolution_falls_through_to_the_first_style_that_exists():
    from services.doc_gen.docx_renderer import _resolve_style

    document = SimpleNamespace(styles=_PartialTemplate({"Heading 6", "Heading 5"}))
    assert _resolve_style(document, ["Heading 6", "Heading 5", "Heading 4"]) == "Heading 4"


def test_style_resolution_reports_nothing_rather_than_raising():
    from services.doc_gen.docx_renderer import _resolve_style

    document = SimpleNamespace(styles=_PartialTemplate({"Quote"}))
    assert _resolve_style(document, ["Quote"]) is None


def _without(prefix, monkeypatch):
    """Render as though the template defined no style starting with ``prefix``."""
    from services.doc_gen import docx_renderer

    real = docx_renderer._resolve_style

    def _filtered(document, candidates):
        return real(document, [c for c in candidates if not c.startswith(prefix)])

    monkeypatch.setattr(docx_renderer, "_resolve_style", _filtered)


def test_a_heading_with_no_style_available_degrades_to_bold_text(monkeypatch):
    """Bold is what is left of "this is a heading" once the style is gone. The
    alternative is a KeyError surfacing as a 500 on the download."""
    _without("Heading", monkeypatch)
    document = _open(render_docx("## Section Title\n\nBody.\n"))
    assert _style_of(document, "Section Title") == "Normal"
    assert _run_named(document, "Section Title").bold is True


def test_a_list_with_no_style_available_falls_back_to_indentation(monkeypatch):
    """The item loses its bullet glyph, which is ugly. It keeps its text and its
    nesting, which is what matters."""
    _without("List", monkeypatch)
    document = _open(render_docx("- outer\n    - inner\n"))
    outer = next(p for p in document.paragraphs if p.text == "outer")
    inner = next(p for p in document.paragraphs if p.text == "inner")
    assert outer.style.name == "Normal"
    assert (inner.paragraph_format.left_indent or 0) > (
        outer.paragraph_format.left_indent or 0
    )


def test_a_table_with_no_bordered_style_still_renders_its_cells(monkeypatch):
    _without("Table", monkeypatch)
    table = _open(render_docx(TABLE_MD)).tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Control", "Status", "Owner"]


def test_a_blockquote_does_not_stamp_the_quote_style_over_a_list():
    """A blockquote can hold a list, and overwriting the paragraph style to mark
    the quotation would take the bullet with it."""
    document = _open(render_docx("> - quoted item\n"))
    quoted = next(p for p in document.paragraphs if p.text == "quoted item")
    assert quoted.style.name == "List Bullet"
    assert (quoted.paragraph_format.left_indent or 0) > 0


def test_text_after_a_nested_list_is_not_lost():
    """The tail of a nested list belongs after it, not back inside the bullet
    that opened the item. Either way it must survive into the document."""
    document = _open(render_docx("- outer\n    - inner\n\n  trailing prose\n"))
    assert "trailing prose" in _all_text(document)
    assert "inner" in _all_text(document)
